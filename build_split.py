# -*- coding: utf-8 -*-
"""Produce journal-required split files: Submission_TitlePage.docx and Submission_MainText.docx (blinded).
Reuses cleaning/parsing helpers from build_submission.py.
"""
import build_submission as B
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def style_doc(doc):
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"; st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

def body_loop(doc, lines, start):
    """Parse body lines (headings/tables/lists/paragraphs) from index `start`."""
    n = len(lines); i = start
    while i < n:
        s = lines[i].rstrip()
        if s.strip() == "" or s.strip() == "---" or s.strip() == "* * *":
            i += 1; continue
        if s.startswith("# "):
            i += 1; continue
        if s.startswith("## "):
            B.add_heading(doc, s[3:].strip(), 1); i += 1; continue
        if s.startswith("### "):
            B.add_heading(doc, s[4:].strip(), 2); i += 1; continue
        if s.strip().startswith("|"):
            block = []
            while i < n and lines[i].strip().startswith("|"):
                block.append(lines[i].rstrip()); i += 1
            B.add_table(doc, block); continue
        if s.startswith("- ") or s.startswith("  - ") or s.startswith("    - "):
            para = doc.add_paragraph(style="List Bullet")
            B.parse_inline(para, s.strip()[2:].strip()); i += 1; continue
        para = doc.add_paragraph(); B.parse_inline(para, s.strip()); i += 1

def build_title_page():
    text, da = B.clean(B.read_text())
    doc = Document(); style_doc(doc)
    title = text.split("\n")[0].lstrip("# ").strip()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title); r.bold = True; r.font.size = Pt(15)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(B.AUTHORS).font.size = Pt(11)
    for a in (B.AFFIL1, B.AFFIL2):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(a).font.size = Pt(9)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(B.EQUAL).font.size = Pt(9); p.runs[0].italic = True
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(B.CORR).font.size = Pt(9)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(B.FUND_COI).font.size = Pt(9)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Running title: BMD and OA genetics: shared architecture, divergent fracture causality").font.size = Pt(9)
    p.runs[0].italic = True
    doc.save("submission_package/Submission_TitlePage.docx")
    print("Saved submission_package/Submission_TitlePage.docx")

def build_main():
    text, da = B.clean(B.read_text())
    lines = text.split("\n")
    doc = Document(); style_doc(doc)
    title = lines[0].lstrip("# ").strip()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title); r.bold = True; r.font.size = Pt(15)
    body_loop(doc, lines, start=1)
    if da:
        doc.add_paragraph()
        h = doc.add_heading(level=2); h.add_run("Data availability")
        p = doc.add_paragraph(); p.add_run(da)
    doc.save("submission_package/Submission_MainText.docx")
    print("Saved submission_package/Submission_MainText.docx")

if __name__ == "__main__":
    build_title_page(); build_main()
