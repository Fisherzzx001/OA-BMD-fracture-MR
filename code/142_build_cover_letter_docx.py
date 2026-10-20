# -*- coding: utf-8 -*-
"""Convert docs/cover_letter_OAC.md -> submission_package/Cover_Letter.docx.

Replaces the stale build_cover_letter.py (which hardcoded an obsolete title and
lacked the OAC S1.1.5 MR block). The markdown is now the canonical source.
"""
import os
import re
from docx import Document
from docx.shared import Pt

PROJ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(PROJ, "docs", "cover_letter_OAC.md")
OUT = os.path.join(PROJ, "submission_package", "Cover_Letter.docx")


def add_runs(p, text):
    """Render inline **bold**, *italic*, `code`."""
    # split on the three inline markers, preserving them
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text)
    for seg in parts:
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**"):
            r = p.add_run(seg[2:-2]); r.bold = True
        elif seg.startswith("`") and seg.endswith("`"):
            r = p.add_run(seg[1:-1]); r.font.name = "Consolas"
        elif seg.startswith("*") and seg.endswith("*"):
            r = p.add_run(seg[1:-1]); r.italic = True
        else:
            p.add_run(seg)


def main():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)

    with open(SRC, encoding="utf-8") as f:
        lines = f.read().splitlines()

    buf = []  # paragraph buffer for consecutive non-list lines

    def flush():
        if not buf:
            return
        text = " ".join(x.strip() for x in buf if x.strip())
        if text:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            add_runs(p, text)
        buf.clear()

    for line in lines:
        if line.startswith("# "):
            flush()
            p = doc.add_heading(level=1)
            add_runs(p, line[2:])
        elif line.startswith("## "):
            flush()
            p = doc.add_heading(level=2)
            add_runs(p, line[3:])
        elif re.match(r"^\s*-\s+", line):
            flush()
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, re.sub(r"^\s*-\s+", "", line))
        elif re.match(r"^\s*\d+\.\s+", line):
            flush()
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\s*\d+\.\s+", "", line))
        elif line.strip() == "":
            flush()
        else:
            buf.append(line)

    flush()
    doc.save(OUT)
    print("Saved", OUT)


if __name__ == "__main__":
    main()
